"""Tests untuk upload Excel: .xlsx, .xls (legacy), multi-sheet, dan non-transaksi."""

import os
import tempfile
from pathlib import Path

import pytest

from app.database.models import (
    Akun,
    JurnalUmum,
    KategoriAkun,
    RoleUser,
    SaldoNormal,
    StatusUpload,
    UploadedFile,
    User,
)
from app.services.ingestion.file_loader import load_file
from app.services.ingestion.ingestion_pipeline import process_uploaded_file


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _write_xlsx(path, sheets: dict) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=name)
        first = False
        for r in rows:
            ws.append(r)
    wb.save(path)


def _write_xls(path, rows) -> None:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Transaksi")
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            ws.write(i, j, val)
    wb.save(path)


def _seed_coa(db) -> None:
    """Seed COA lengkap dari knowledge/datasets/coa.csv supaya keyword
    mapping (csv_to_jurnal) bisa memetakan deskripsi ke kode akun dengan benar."""
    import csv

    from app.config.settings import settings

    coa_path = Path(settings.KNOWLEDGE_DIR) / "datasets" / "coa.csv"
    with open(coa_path, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            db.add(Akun(
                kode_akun=row["kode_akun"].strip(),
                nama_akun=row["nama_akun"].strip(),
                kategori=KategoriAkun(row["kategori"].strip()),
                sub_kategori=row.get("sub_kategori") or None,
                saldo_normal=SaldoNormal(row["saldo_normal"].strip()),
                is_active=True,
            ))
    db.commit()


def _make_user(db) -> User:
    user = User(
        email="excel_test@example.com",
        hashed_password="x",
        full_name="Excel Test",
        company_name="T",
        role=RoleUser.OWNER,
        is_active=True,
        email_verified=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@pytest.fixture()
def excel_user(db):
    _seed_coa(db)
    return _make_user(db)


def _make_uploaded_file(db, user, path: str, file_type: str) -> UploadedFile:
    uf = UploadedFile(
        original_filename=os.path.basename(path),
        stored_path=path,
        file_type=file_type,
        file_size_bytes=os.path.getsize(path),
        uploaded_by_id=user.id,
    )
    db.add(uf)
    db.commit()
    db.refresh(uf)
    return uf


# ---------------------------------------------------------------------------
# Loader level — format & multi-sheet detection
# ---------------------------------------------------------------------------
class TestExcelLoader:
    def test_xlsx_single_sheet(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            _write_xlsx(path, {"Transaksi": [
                ["tanggal", "deskripsi", "jumlah"],
                ["2025-01-01", "Penjualan", 100000],
                ["2025-01-02", "Beli bahan", 50000],
            ]})
            doc = load_file(path, "xlsx")
            assert doc.file_type == "xlsx"
            df = doc.get_transaction_dataframe()
            assert df is not None
            assert "tanggal" in df.columns
            assert len(df) == 2
        finally:
            os.unlink(path)

    def test_xlsx_multisheet_picks_second_sheet(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            _write_xlsx(path, {
                "Cover": [["Judul Laporan"], ["Periode 2025"]],
                "Transaksi": [
                    ["tanggal", "deskripsi", "jumlah"],
                    ["2025-01-01", "Penjualan", 100000],
                ],
            })
            doc = load_file(path, "xlsx")
            df = doc.get_transaction_dataframe()
            # Sheet terpilih harus "Transaksi" (ada kolom tanggal+nominal),
            # bukan "Cover" yang hanya berisi teks.
            assert df is not None
            assert list(df["deskripsi"]) == ["Penjualan"]
        finally:
            os.unlink(path)

    def test_xls_legacy(self):
        pytest.importorskip("xlwt")
        with tempfile.NamedTemporaryFile(suffix=".xls", delete=False) as f:
            path = f.name
        try:
            _write_xls(path, [
                ["tanggal", "deskripsi", "jumlah"],
                ["2025-01-01", "Penjualan", 100000],
            ])
            # Validator memetakan .xls -> file_type "xlsx"; engine dipilih
            # berdasarkan ekstensi file (.xls -> xlrd).
            doc = load_file(path, "xlsx")
            df = doc.get_transaction_dataframe()
            assert df is not None
            assert len(df) == 1
        finally:
            os.unlink(path)

    def test_nontransaction_workbook_has_no_sheet(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            _write_xlsx(path, {"Notes": [
                ["Catatan"],
                ["Ini bukan data transaksi"],
            ]})
            doc = load_file(path, "xlsx")
            assert doc.get_transaction_dataframe() is None
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Service level — end-to-end melalui process_uploaded_file
# ---------------------------------------------------------------------------
class TestExcelIngestion:
    def test_multisheet_xlsx_creates_jurnal(self, db, excel_user):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            _write_xlsx(path, {
                "Cover": [["Laporan Keuangan 2025"]],
                "Transaksi": [
                    ["tanggal", "deskripsi", "jumlah", "jenis", "kategori"],
                    ["2025-01-01", "Penjualan barang", 100000, "masuk", "pendapatan"],
                    ["2025-01-02", "Bayar listrik", 50000, "keluar", "beban_operasional"],
                ],
            })
            uf = _make_uploaded_file(db, excel_user, path, "xlsx")
            process_uploaded_file(db, uf)
            db.commit()
            assert uf.status == StatusUpload.POSTED
            jurnals = db.query(JurnalUmum).filter(JurnalUmum.sumber_upload_id == uf.id).all()
            assert len(jurnals) == 2
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_xls_legacy_creates_jurnal(self, db, excel_user):
        pytest.importorskip("xlwt")
        with tempfile.NamedTemporaryFile(suffix=".xls", delete=False) as f:
            path = f.name
        try:
            _write_xls(path, [
                ["tanggal", "deskripsi", "jumlah", "jenis", "kategori"],
                ["2025-01-01", "Penjualan barang", 100000, "masuk", "pendapatan"],
            ])
            uf = _make_uploaded_file(db, excel_user, path, "xlsx")  # .xls -> "xlsx"
            process_uploaded_file(db, uf)
            db.commit()
            assert uf.status == StatusUpload.POSTED
            jurnals = db.query(JurnalUmum).filter(JurnalUmum.sumber_upload_id == uf.id).all()
            assert len(jurnals) == 1
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_nontransaction_xlsx_no_jurnal(self, db, excel_user, monkeypatch):
        # Cegah embedding/Qdrant (tidak tersedia di test env) lewat RAG pipeline.
        monkeypatch.setattr(
            "app.services.ingestion.ingestion_pipeline.embed_chunks",
            lambda texts: [[0.0] * 384 for _ in texts],
        )
        monkeypatch.setattr(
            "app.services.ingestion.ingestion_pipeline.qdrant_service.upsert_chunks",
            lambda vectors, payloads: [f"id-{i}" for i in range(len(payloads))],
        )
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            _write_xlsx(path, {"Notes": [
                ["Catatan"],
                ["Ini bukan data transaksi"],
            ]})
            uf = _make_uploaded_file(db, excel_user, path, "xlsx")
            process_uploaded_file(db, uf)
            db.commit()
            jurnals = db.query(JurnalUmum).filter(JurnalUmum.sumber_upload_id == uf.id).all()
            assert len(jurnals) == 0
            # Tak crash; status INGESTED (RAG jalan) atau FAILED (bila embedding
            # tetap gagal) — yang penting tak ada jurnal terbuat.
            assert uf.status in (StatusUpload.INGESTED, StatusUpload.FAILED)
        finally:
            if os.path.exists(path):
                os.unlink(path)
